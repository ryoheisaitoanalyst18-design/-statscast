#!/usr/bin/env python3
"""ローカル取込アプリ — TrackMan CSV/xlsx/zip をブラウザから投入して本番反映まで自動実行する。

    python3 scripts/ingest_server.py        # → http://127.0.0.1:8787

ブラウザで CSV/xlsx/zip をドロップ → 「取り込んで本番に反映」→ 進捗ログが流れ、
検証ゲート (validate_data.py) が通ったときだけ push される。

やっていること (中身は既存の update_and_deploy.sh に丸投げ):
  1. アップロードを ~/statscast_inbox/<日時>/ に保存
  2. zip は展開し、各ファイルを tokyo-baseball/trackman_io.py で判定して読む
     - 投球データ → 「列の和集合」で 1 本に結合 (列ズレ破損行を作らない)
     - ポジショニング (守備位置) → ~/ubuntu_data/positioning/ へ保存 (守備分析ページの入力)
     - ヘッダー無し・集計表など投球データでないもの → 理由をログに出して取り込まない
       (旧実装は何でも結合したため、ヘッダー無し CSV の値が列名化した破損行がマスターに入った)
  3. update_and_deploy.sh <combined.csv> --clean を実行
     → マスターCSVバックアップ → QA(打席合体/六大学外/破損行) → --clean 除外 →
       試合単位の置き換え (同じ試合の入れ直しは二重にせず置き換え) → マージ → 全再生成 →
       検証ゲート → push → 本番URL確認
  4. 出力を 1 行ずつブラウザへ中継

localhost からしか接続を受けない。ファイル判定だけ pandas (パイプラインと同じ環境) を使う。
"""
from __future__ import annotations

import csv
import json
import os
import shutil
import subprocess
import sys
import threading
import time
import urllib.parse
import zipfile
from datetime import datetime
from http.server import BaseHTTPRequestHandler, ThreadingHTTPServer
from pathlib import Path

REPO_DIR = Path(__file__).resolve().parent.parent
DEPLOY_SH = REPO_DIR / "update_and_deploy.sh"
INBOX = Path.home() / "statscast_inbox"
PIPELINE_DIR = Path.home() / "tokyo-baseball"
POSITIONING_DIR = Path(os.environ.get("STATSCAST_POSITIONING", Path.home() / "ubuntu_data" / "positioning"))
HOST = "127.0.0.1"
PORT = int(os.environ.get("INGEST_PORT", "8787"))
MAX_UPLOAD = 2 * 1024 * 1024 * 1024  # 1ファイル 2GB まで

csv.field_size_limit(10 * 1024 * 1024)


# ============================================================
# ジョブ (同時に 1 本だけ)
# ============================================================
class Job:
    def __init__(self, sid: str) -> None:
        self.sid = sid
        self.dir = INBOX / sid
        self.dir.mkdir(parents=True, exist_ok=True)
        self.files: list[Path] = []
        self.lines: list[str] = []
        self.lock = threading.Lock()
        self.state = "idle"  # idle | running | success | failed
        self.started_at: float | None = None
        self.finished_at: float | None = None
        self.proc: subprocess.Popen | None = None
        # 編集モード: /prepare で結合済み combined.csv を作り、メモリ上の表を編集してから /start
        self.prepared = False
        self.df = None          # pandas.DataFrame (全列 str)
        self.flags: dict[int, str] = {}
        self.edit_lock = threading.Lock()

    def log(self, text: str) -> None:
        with self.lock:
            self.lines.extend(text.rstrip("\n").split("\n"))

    def tail(self, offset: int) -> tuple[list[str], int]:
        with self.lock:
            return self.lines[offset:], len(self.lines)

    def snapshot(self, offset: int) -> dict:
        new, total = self.tail(offset)
        elapsed = 0.0
        if self.started_at:
            elapsed = (self.finished_at or time.time()) - self.started_at
        return {
            "sid": self.sid,
            "state": self.state,
            "lines": new,
            "offset": total,
            "elapsed": round(elapsed),
            "files": [f.name for f in self.files],
        }


JOBS: dict[str, Job] = {}
CURRENT: Job | None = None
CURRENT_LOCK = threading.Lock()


# ============================================================
# zip 展開 + CSV 結合
# ============================================================
def unique_path(directory: Path, name: str) -> Path:
    """directory/name が既にあれば name_2.csv, name_3.csv … と衝突しない名前を返す。

    TrackMan の zip は日付フォルダごとに同じファイル名 (例 game.csv) を持つことがあり、
    basename だけで展開すると後の 1 本が前の 1 本を上書きして試合が丸ごと消える。
    """
    candidate = directory / name
    if not candidate.exists():
        return candidate
    stem, suffix = candidate.stem, candidate.suffix
    for i in range(2, 10000):
        candidate = directory / f"{stem}_{i}{suffix}"
        if not candidate.exists():
            return candidate
    raise RuntimeError(f"名前の衝突を解決できません: {directory / name}")


def xlsx_to_csv(src: Path, job: Job) -> Path:
    """TrackMan の .xlsx 納品を CSV に変換して返す (2026秋シーズンから .xlsx 納品あり)。

    値のみを読む (read_only=True, data_only=True)。日時セルは ISO 形式の文字列にして、
    CSV 経由でも UTCDate/Date の解釈がブレないようにする。
    """
    import openpyxl  # 変換時のみ必要 (CSV/zip だけなら stdlib のまま動く)
    from datetime import date as _date

    def cell(v):
        # xlsx は整数列も float で返す。マスターCSV は PitchNo=1 / Inning=1 の形なので合わせる。
        if isinstance(v, float) and v.is_integer():
            return int(v)
        # 日付セルは 00:00:00 が付く。既存行は 'YYYY-MM-DD' なので日付だけにする。
        if isinstance(v, datetime):
            return v.date().isoformat() if v.time() == v.time().min else v.isoformat(sep=" ")
        if isinstance(v, _date):
            return v.isoformat()
        return "" if v is None else v

    out = src.with_suffix(".csv")
    out = unique_path(out.parent, out.name)
    wb = openpyxl.load_workbook(src, read_only=True, data_only=True)
    ws = wb[wb.sheetnames[0]]
    n = 0
    with open(out, "w", newline="", encoding="utf-8") as fh:
        writer = csv.writer(fh)
        for row in ws.iter_rows(values_only=True):
            if all(v is None for v in row):
                continue
            writer.writerow([cell(v) for v in row])
            n += 1
    wb.close()
    job.log(f"  xlsx 変換: {src.name} → {out.name} ({max(n - 1, 0)} 行)")
    return out


def collect_csvs(job: Job) -> list[Path]:
    """アップロードされた CSV/xlsx と、zip 内のそれらを集める。"""
    found: list[Path] = []

    def extract_zip(zpath: Path, depth: int = 0) -> None:
        dest = job.dir / f"{zpath.stem}_extracted"
        dest.mkdir(exist_ok=True)
        n_from_zip = 0
        with zipfile.ZipFile(zpath) as z:
            for info in z.infolist():
                member = info.filename
                # 日本語ファイル名: UTF-8 フラグの無い zip は cp932 で名前を復元する
                if not (info.flag_bits & 0x800):
                    try:
                        member = member.encode("cp437").decode("cp932")
                    except (UnicodeEncodeError, UnicodeDecodeError):
                        pass
                low = member.lower()
                if member.endswith("/") or "__macosx" in low or not low.endswith((".csv", ".xlsx", ".zip")):
                    continue
                # zip slip 対策: 展開先を dest 配下に強制する。
                # さらに同名 member の上書き (試合の消失) を防ぐため一意名にする。
                safe = unique_path(dest, Path(member).name)
                with z.open(info) as src, open(safe, "wb") as out:
                    shutil.copyfileobj(src, out)
                if safe.name != Path(member).name:
                    job.log(f"    名前衝突を回避: {member} → {safe.name}")
                if safe.suffix.lower() == ".zip":
                    # zip の中の zip (フレッシュトーナメント分などで実例あり) も展開する
                    if depth < 3:
                        extract_zip(safe, depth + 1)
                    continue
                found.append(xlsx_to_csv(safe, job) if safe.suffix.lower() == ".xlsx" else safe)
                n_from_zip += 1
        job.log(f"  zip 展開: {zpath.name} → ファイル {n_from_zip} 本")

    for f in job.files:
        if f.suffix.lower() == ".zip":
            extract_zip(f)
        elif f.suffix.lower() == ".csv":
            found.append(f)
        elif f.suffix.lower() == ".xlsx":
            found.append(xlsx_to_csv(f, job))
        else:
            job.log(f"  スキップ (CSV/xlsx/zip ではない): {f.name}")
    return sorted(found)


def combine_csvs(csvs: list[Path], out_path: Path, job: Job) -> int:
    """投球データのファイルだけを「列の和集合」で 1 本に結合する。

    各ファイルは trackman_io.read_trackman_file で判定する (行末カンマ・タイトル行・
    列名誤字を吸収)。ポジショニングは守備分析の置き場へ保存し、投球データでないものは
    理由をログに出して結合しない。列構成が違うファイルは欠けた列を空欄で埋める
    (単純連結すると列ズレ破損行が生まれる。過去に 1,671 球の実害)。
    """
    sys.path.insert(0, str(PIPELINE_DIR))
    import hashlib
    import pandas as pd
    import trackman_io

    frames = []
    for p in csvs:
        df, kind, reason = trackman_io.read_trackman_file(str(p))
        if kind == "pitch":
            frames.append(df)
            job.log(f"    {p.name}: {len(df)} 球")
        elif kind == "positioning":
            POSITIONING_DIR.mkdir(parents=True, exist_ok=True)
            digest = hashlib.md5(p.read_bytes()).hexdigest()[:10]
            dest = POSITIONING_DIR / f"{p.stem}__{digest}.csv"
            if not dest.exists():
                df.to_csv(dest, index=False)
            job.log(f"    {p.name}: 守備位置データ → {dest.parent.name}/ に保存 (守備分析で使用)")
        else:
            job.log(f"    ⚠ {p.name}: 取り込みません — {reason}")
    if not frames:
        return 0
    combined = pd.concat(frames, ignore_index=True, sort=False)
    combined.to_csv(out_path, index=False)
    return len(combined)


# ============================================================
# 編集 (取込前に結合済み CSV を表で直す)
# ============================================================
# 表の既定表示列 (存在するものだけ)。「全列」トグルで全列表示。
KEY_COLUMNS = ["PitchNo", "Date", "Time", "GameID", "Inning", "Top/Bottom", "Outs",
               "Balls", "Strikes", "PAofInning", "PitchofPA", "Pitcher", "PitcherTeam",
               "Batter", "BatterTeam", "TaggedPitchType", "PitchCall", "KorBB",
               "TaggedHitType", "PlayResult", "OutsOnPlay", "RunsScored",
               "RelSpeed", "SpinRate", "InducedVertBreak", "HorzBreak",
               "PlateLocHeight", "PlateLocSide", "ExitSpeed", "Angle"]


def compute_flags(df) -> dict[int, str]:
    """QA で問題になる行に理由を付ける (trackman_qa と同じ判定)。"""
    sys.path.insert(0, str(PIPELINE_DIR))
    import pandas as pd
    import trackman_qa as qa

    flags: dict[int, list[str]] = {}
    broken = qa.check_broken_pa(df.replace("", pd.NA))
    if not broken.empty:
        keys = set(map(tuple, broken[qa.PA_KEYS].astype(str).values))
        pa = df[qa.PA_KEYS].astype(str).agg(tuple, axis=1)
        for i in df.index[pa.isin(keys)]:
            flags.setdefault(int(i), []).append("打席合体")
    if {"BatterTeam", "PitcherTeam"}.issubset(df.columns):
        bt = qa._norm_team(df["BatterTeam"], qa.DEFAULT_TEAM_MAP)
        pt = qa._norm_team(df["PitcherTeam"], qa.DEFAULT_TEAM_MAP)
        non_six = ~(bt.isin(qa.DEFAULT_VALID_TEAMS) | pt.isin(qa.DEFAULT_VALID_TEAMS))
        suspect = non_six & (qa._is_suspect_team(df["BatterTeam"])
                             | qa._is_suspect_team(df["PitcherTeam"]))
        for i in df.index[suspect]:
            flags.setdefault(int(i), []).append("破損/列ズレ疑い")
        for i in df.index[non_six & ~suspect]:
            flags.setdefault(int(i), []).append("六大学外")
    return {k: "・".join(v) for k, v in flags.items()}


def flag_summary(job: Job) -> dict:
    counts: dict[str, int] = {}
    for reason in job.flags.values():
        for r in reason.split("・"):
            counts[r] = counts.get(r, 0) + 1
    games = int(job.df["GameID"].nunique()) if "GameID" in job.df.columns else 0
    return {"rows": len(job.df), "cols": len(job.df.columns), "games": games,
            "issues": counts, "flagged": len(job.flags)}


def load_for_edit(job: Job) -> None:
    import pandas as pd
    # 全列文字列で読む (編集しないセルの表記を一切変えない)
    df = pd.read_csv(job.dir / "combined.csv", dtype=str, keep_default_na=False)
    job.df = df
    job.flags = compute_flags(df)


def save_edited(job: Job, note: dict) -> None:
    """編集後の表を combined.csv に書き戻し、操作ログを edits.jsonl に残す。"""
    job.df = job.df.reset_index(drop=True)
    job.df.to_csv(job.dir / "combined.csv", index=False)
    job.flags = compute_flags(job.df)
    with open(job.dir / "edits.jsonl", "a", encoding="utf-8") as fh:
        fh.write(json.dumps({"at": datetime.now().isoformat(timespec="seconds"), **note},
                            ensure_ascii=False) + "\n")


def filtered_index(job: Job, q: dict[str, str]):
    df = job.df
    mask = None
    col, val = q.get("fcol", ""), q.get("fval", "")
    if val:
        cols = [col] if col in df.columns else [c for c in df.columns if c in KEY_COLUMNS]
        m = None
        for c in cols:
            hit = df[c].str.contains(val, case=False, regex=False)
            m = hit if m is None else (m | hit)
        mask = m
    idx = df.index if mask is None else df.index[mask]
    if q.get("issues") == "1":
        idx = idx[idx.isin(list(job.flags))]
    return idx


# ============================================================
# 実行
# ============================================================
def run_job(job: Job) -> None:
    global CURRENT
    job.state = "running"
    job.started_at = time.time()
    try:
        combined = job.dir / "combined.csv"
        if job.prepared:
            rows = len(job.df)
            if rows == 0:
                job.log("エラー: 編集で全行が削除されています。取り込むデータがありません。")
                job.state = "failed"
                return
            job.log("━━━ 編集済みの結合 CSV を使います ━━━")
            if (job.dir / "edits.jsonl").exists():
                n = sum(1 for _ in open(job.dir / "edits.jsonl", encoding="utf-8"))
                job.log(f"  編集操作 {n} 件 (記録: {job.dir / 'edits.jsonl'})")
        else:
            job.log("━━━ 入力ファイルの準備 ━━━")
            csvs = collect_csvs(job)
            if not csvs:
                job.log("エラー: CSV が 1 本も見つかりませんでした。")
                job.state = "failed"
                return
            job.log(f"CSV {len(csvs)} 本を結合します:")
            rows = combine_csvs(csvs, combined, job)
        if rows > 0:
            job.log(f"結合完了: {rows} 行 → {combined}")
            args = [str(combined), "--clean"]
        else:
            # 投球データが無い (守備位置データだけ等) → マスターは変えずに全再生成だけ行う
            job.log("投球データのファイルはありませんでした。マスターは変更せず全再生成します。")
            args = ["--regenerate"]
        job.log("")
        job.log("━━━ update_and_deploy.sh 開始 (再生成に 5〜9 分かかります) ━━━")

        env = dict(os.environ, PYTHONUNBUFFERED="1")
        proc = subprocess.Popen(
            ["bash", str(DEPLOY_SH), *args],
            cwd=str(REPO_DIR),
            stdout=subprocess.PIPE,
            stderr=subprocess.STDOUT,
            text=True,
            bufsize=1,
            env=env,
        )
        job.proc = proc
        assert proc.stdout is not None
        for line in proc.stdout:
            job.log(line)
        code = proc.wait()
        job.log("")
        if code == 0:
            job.log("✅ 完了: 本番サイトに反映されました。")
            job.state = "success"
        else:
            job.log(f"❌ 失敗 (exit {code})。上のログを確認してください。")
            job.log("   マスターCSVは実行前にバックアップ済みです "
                    "(~/ubuntu_data/trackman_data.backup_*.csv)。")
            job.state = "failed"
    except Exception as exc:  # noqa: BLE001 — 何が起きてもブラウザに理由を出す
        job.log(f"❌ 例外: {type(exc).__name__}: {exc}")
        job.state = "failed"
    finally:
        job.finished_at = time.time()
        with CURRENT_LOCK:
            if CURRENT is job:
                CURRENT = None


# ============================================================
# HTTP
# ============================================================
PAGE = """<!doctype html>
<html lang="ja"><head><meta charset="utf-8">
<meta name="viewport" content="width=device-width,initial-scale=1">
<title>statscast 取込</title>
<style>
:root{color-scheme:light dark;--bg:#fff;--fg:#111;--mut:#666;--line:#ddd;--card:#fafafa;--acc:#1f6feb}
@media(prefers-color-scheme:dark){:root{--bg:#0d1117;--fg:#e6edf3;--mut:#8b949e;--line:#30363d;--card:#161b22;--acc:#4493f8}}
*{box-sizing:border-box}
body{margin:0;background:var(--bg);color:var(--fg);font:15px/1.6 system-ui,"Hiragino Sans","Noto Sans JP",sans-serif}
.wrap{max-width:1200px;margin:0 auto;padding:32px 20px 64px}
h1{font-size:20px;margin:0 0 4px}
.sub{color:var(--mut);font-size:13px;margin-bottom:24px}
#drop{border:2px dashed var(--line);border-radius:12px;padding:44px 20px;text-align:center;
  background:var(--card);cursor:pointer;transition:.15s}
#drop:hover,#drop.over{border-color:var(--acc);background:color-mix(in srgb,var(--acc) 8%,var(--card))}
#drop b{display:block;font-size:16px;margin-bottom:6px}
#drop span{color:var(--mut);font-size:13px}
ul{list-style:none;padding:0;margin:16px 0 0}
li{display:flex;justify-content:space-between;gap:12px;padding:8px 12px;background:var(--card);
  border:1px solid var(--line);border-radius:8px;margin-bottom:6px;font-size:13px}
li .sz{color:var(--mut);white-space:nowrap}
.row{display:flex;gap:12px;align-items:center;margin-top:20px;flex-wrap:wrap}
button{font:inherit;font-weight:600;padding:10px 20px;border-radius:8px;border:1px solid transparent;
  background:var(--acc);color:#fff;cursor:pointer}
button:disabled{opacity:.45;cursor:not-allowed}
button.ghost{background:transparent;color:var(--fg);border-color:var(--line);font-weight:400}
.badge{font-size:12px;padding:4px 10px;border-radius:99px;border:1px solid var(--line);color:var(--mut)}
.badge.running{border-color:var(--acc);color:var(--acc)}
.badge.success{border-color:#2ea043;color:#2ea043}
.badge.failed{border-color:#f85149;color:#f85149}
pre{margin:16px 0 0;padding:16px;background:var(--card);border:1px solid var(--line);border-radius:10px;
  max-height:460px;overflow:auto;font:12px/1.55 ui-monospace,SFMono-Regular,Menlo,monospace;
  white-space:pre-wrap;word-break:break-word}
.hide{display:none}
#ed{margin-top:24px}
.sum{display:flex;gap:8px;flex-wrap:wrap;margin-bottom:12px}
.sum .badge.warn{border-color:#d29922;color:#d29922}
.tools{display:flex;gap:8px;flex-wrap:wrap;align-items:center;margin:8px 0;font-size:13px}
.tools input,.tools select{font:inherit;padding:6px 8px;border-radius:6px;border:1px solid var(--line);
  background:var(--bg);color:var(--fg)}
.tools input[type=text]{width:130px}
.tools button{padding:6px 12px;font-size:13px}
.tools label{display:flex;gap:4px;align-items:center;color:var(--mut)}
.tbl{overflow:auto;max-height:560px;border:1px solid var(--line);border-radius:10px;margin-top:8px}
table{border-collapse:collapse;font:12px/1.4 ui-monospace,SFMono-Regular,Menlo,monospace;white-space:nowrap}
th,td{border-bottom:1px solid var(--line);border-right:1px solid var(--line);padding:3px 6px}
th{position:sticky;top:0;background:var(--card);z-index:1;text-align:left}
td[contenteditable]{min-width:40px;outline:none}
td[contenteditable]:focus{box-shadow:inset 0 0 0 2px var(--acc)}
td.dirty{background:color-mix(in srgb,#d29922 22%,transparent)}
tr.flag td:first-child{border-left:3px solid #f85149}
tr.del td{text-decoration:line-through;opacity:.45}
td.meta{color:var(--mut);font-size:11px}
.pager{display:flex;gap:8px;align-items:center;margin-top:8px;font-size:13px;color:var(--mut)}
.pager button{padding:4px 10px;font-size:13px}
</style></head><body><div class="wrap">
<h1>試合データ取込</h1>
<div class="sub">TrackMan の CSV / zip を置くと、QA → マスター取込 → 全再生成 → 検証 → 本番反映まで自動で行います。</div>

<div id="drop">
  <b>ここに CSV / zip をドロップ</b>
  <span>クリックして選択もできます（複数可・zip の中身も自動で取り出します）</span>
</div>
<input type="file" id="pick" multiple accept=".csv,.xlsx,.zip" class="hide">
<ul id="list"></ul>

<div class="row">
  <button id="edit" class="ghost" disabled>プレビュー・編集</button>
  <button id="go" disabled>取り込んで本番に反映</button>
  <button id="clear" class="ghost">選び直す</button>
  <span id="badge" class="badge">待機中</span>
  <span id="time" class="badge hide"></span>
</div>

<section id="ed" class="hide">
  <div class="sum" id="sum"></div>
  <div class="tools">
    <label>絞り込み <select id="fcol"><option value="">主要列すべて</option></select></label>
    <input type="text" id="fval" placeholder="含む文字">
    <label><input type="checkbox" id="issues"> 問題行のみ</label>
    <label><input type="checkbox" id="allc"> 全列表示</label>
  </div>
  <div class="tools">
    <label>一括置換 <select id="rcol"></select></label>
    <input type="text" id="rfind" placeholder="検索">→<input type="text" id="rrepl" placeholder="置換後">
    <label><input type="checkbox" id="rwhole"> セル完全一致</label>
    <button id="rgo" class="ghost">絞り込み中の行に置換</button>
  </div>
  <div class="tools">
    <button id="save">変更を保存</button>
    <button id="discard" class="ghost">未保存の変更を破棄</button>
    <button id="revert" class="ghost">元ファイルに戻す</button>
    <span id="dirty" class="badge">未保存 0</span>
  </div>
  <div class="tbl"><table id="tb"></table></div>
  <div class="pager"><button id="prev" class="ghost">‹ 前</button><span id="pinfo"></span>
    <button id="next" class="ghost">次 ›</button>
    <span>セルをクリックで編集・Enterで確定 / 左端の × で行削除 / 赤線=QAで問題のある行</span></div>
</section>

<pre id="log" class="hide"></pre>
</div><script>
let files=[],sid=null,offset=0,timer=null;
const $=i=>document.getElementById(i);
const fmt=n=>n>1048576?(n/1048576).toFixed(1)+" MB":(n/1024).toFixed(0)+" KB";
let prepared=false;
const esc=v=>String(v).replace(/[&<>"]/g,c=>({"&":"&amp;","<":"&lt;",">":"&gt;",'"':"&quot;"})[c]);
function render(){
  $("list").innerHTML=files.map(f=>`<li><span>${esc(f.name)}</span><span class="sz">${fmt(f.size)}</span></li>`).join("");
  $("go").disabled=files.length===0; $("edit").disabled=files.length===0||prepared;
}
async function uploadAll(){
  const r=await fetch("/session",{method:"POST"});
  if(!r.ok) throw new Error(await r.text());
  sid=(await r.json()).sid;
  for(const f of files){
    $("log").textContent+=`アップロード中: ${f.name} ...\\n`;
    const u=await fetch(`/upload?sid=${sid}&name=`+encodeURIComponent(f.name),{method:"POST",body:f});
    if(!u.ok) throw new Error(await u.text());
  }
}
async function post(url,body){
  const r=await fetch(url,{method:"POST",headers:{"Content-Type":"application/json"},body:JSON.stringify(body||{})});
  if(!r.ok) throw new Error(await r.text());
  return r.json();
}
/* ---- 編集 ---- */
let page=0,size=100,cur=null,edits=new Map(),dels=new Set();
function showSum(m){
  const iss=Object.entries(m.issues).map(([k,v])=>`<span class="badge warn">${esc(k)} ${v}球</span>`).join("");
  $("sum").innerHTML=`<span class="badge">${m.rows.toLocaleString()} 球</span><span class="badge">${m.games} 試合</span>`+
    `<span class="badge">${m.cols} 列</span>`+(iss||`<span class="badge success">QA 問題なし</span>`);
}
function markDirty(){ $("dirty").textContent=`未保存 ${edits.size+dels.size}`; $("dirty").className="badge"+(edits.size+dels.size?" running":""); }
function q(){ return `fcol=${encodeURIComponent($("fcol").value)}&fval=${encodeURIComponent($("fval").value)}&issues=${$("issues").checked?1:0}`; }
async function load(){
  const r=await fetch(`/table?sid=${sid}&page=${page}&size=${size}&all=${$("allc").checked?1:0}&`+q());
  if(!r.ok){ alert(await r.text()); return; }
  cur=await r.json(); showSum(cur.summary);
  const fc=$("fcol").value, rc=$("rcol").value;
  $("fcol").innerHTML=`<option value="">主要列すべて</option>`+cur.allColumns.map(c=>`<option>${esc(c)}</option>`).join("");
  $("rcol").innerHTML=cur.allColumns.map(c=>`<option>${esc(c)}</option>`).join("");
  $("fcol").value=fc; if(rc) $("rcol").value=rc; else if(cur.allColumns.includes("GameID")) $("rcol").value="GameID";
  let h=`<tr><th></th><th>#</th><th>QA</th>`+cur.columns.map(c=>`<th>${esc(c)}</th>`).join("")+`</tr>`;
  for(const [i,flag,...vals] of cur.rows){
    h+=`<tr data-r="${i}" class="${flag?"flag":""} ${dels.has(i)?"del":""}"><td><button class="ghost x" title="行を削除/取消" style="padding:0 6px">×</button></td>`+
       `<td class="meta">${i}</td><td class="meta">${esc(flag)}</td>`+
       vals.map((v,k)=>{const key=i+"\\t"+cur.columns[k],e=edits.has(key);
         return `<td contenteditable="plaintext-only" spellcheck="false" data-c="${esc(cur.columns[k])}" data-o="${esc(v)}" class="${e?"dirty":""}">${esc(e?edits.get(key):v)}</td>`}).join("")+`</tr>`;
  }
  $("tb").innerHTML=h;
  const last=Math.max(Math.ceil(cur.total/size)-1,0);
  $("pinfo").textContent=`${cur.total.toLocaleString()} 行中 ${cur.total?page*size+1:0}–${Math.min((page+1)*size,cur.total)}`;
  $("prev").disabled=page===0; $("next").disabled=page>=last;
}
$("tb").addEventListener("focusout",e=>{
  const td=e.target; if(td.tagName!=="TD"||!td.dataset.c) return;
  const r=+td.parentElement.dataset.r,key=r+"\\t"+td.dataset.c,v=td.textContent;
  if(v===td.dataset.o) edits.delete(key); else edits.set(key,v);
  td.classList.toggle("dirty",edits.has(key)); markDirty();
});
$("tb").addEventListener("keydown",e=>{ if(e.key==="Enter"&&e.target.tagName==="TD"){e.preventDefault();e.target.blur();} });
$("tb").addEventListener("click",e=>{
  if(!e.target.classList.contains("x")) return;
  const tr=e.target.closest("tr"),r=+tr.dataset.r;
  dels.has(r)?dels.delete(r):dels.add(r); tr.classList.toggle("del",dels.has(r)); markDirty();
});
const confirmLeave=()=>!(edits.size+dels.size)||confirm("未保存の変更があります。保存せずに移動しますか？");
async function saveEdits(){
  if(!(edits.size+dels.size)) return;
  const cells=[...edits].map(([k,v])=>{const t=k.indexOf("\\t");return [+k.slice(0,t),k.slice(t+1),v]});
  const d=await post(`/save?sid=${sid}`,{cells,delete:[...dels]});
  edits.clear(); dels.clear(); markDirty();
  $("log").textContent+=`保存: ${d.changed} セル変更 / ${d.deleted} 行削除\\n`;
}
$("save").onclick=async()=>{ try{ await saveEdits(); await load(); }catch(e){ alert(e.message); } };
$("discard").onclick=()=>{ edits.clear(); dels.clear(); markDirty(); load(); };
$("revert").onclick=async()=>{
  if(!confirm("すべての編集を取り消して、アップロード直後の状態に戻しますか？")) return;
  edits.clear(); dels.clear(); markDirty(); await post(`/revert?sid=${sid}`); page=0; load();
};
$("rgo").onclick=async()=>{
  try{
    await saveEdits();
    const b={col:$("rcol").value,find:$("rfind").value,replace:$("rrepl").value,whole:$("rwhole").checked,
             fcol:$("fcol").value,fval:$("fval").value,issues:$("issues").checked?"1":""};
    if(!confirm(`${b.col} 列の「${b.find}」→「${b.replace}」を、絞り込み中の ${cur.total} 行に適用しますか？`)) return;
    const d=await post(`/replace?sid=${sid}`,b);
    $("log").textContent+=`一括置換: ${b.col} ${d.replaced} 行\\n`; load();
  }catch(e){ alert(e.message); }
};
for(const id of ["fcol","issues","allc"]) $(id).onchange=()=>{ if(confirmLeave()){edits.clear();dels.clear();markDirty();page=0;load();} };
let ft; $("fval").oninput=()=>{ clearTimeout(ft); ft=setTimeout(()=>{ if(confirmLeave()){edits.clear();dels.clear();markDirty();page=0;load();} },350); };
$("prev").onclick=()=>{ if(confirmLeave()){edits.clear();dels.clear();markDirty();page--;load();} };
$("next").onclick=()=>{ if(confirmLeave()){edits.clear();dels.clear();markDirty();page++;load();} };
$("edit").onclick=async()=>{
  $("edit").disabled=true; $("go").disabled=true; $("clear").disabled=true;
  $("log").className=""; $("log").textContent="";
  try{
    await uploadAll();
    $("log").textContent+="結合・QA 中...\\n";
    const d=await post(`/prepare?sid=${sid}`);
    $("log").textContent+=d.log.join("\\n")+"\\n";
    prepared=true; $("ed").className=""; page=0; await load();
  }catch(e){ $("log").textContent+="\\n❌ "+e.message+"\\n"; $("edit").disabled=false; }
  $("go").disabled=false; $("clear").disabled=false;
};
function add(fs){
  for(const f of [...fs].filter(f=>/\\.(csv|xlsx|zip)$/i.test(f.name))){
    if(!files.some(g=>g.name===f.name&&g.size===f.size&&g.lastModified===f.lastModified)) files.push(f);
  }
  render();
}
const drop=$("drop");
drop.onclick=()=>$("pick").click();
$("pick").onchange=e=>add(e.target.files);
drop.ondragover=e=>{e.preventDefault();drop.classList.add("over")};
drop.ondragleave=()=>drop.classList.remove("over");
drop.ondrop=e=>{e.preventDefault();drop.classList.remove("over");add(e.dataTransfer.files)};
$("clear").onclick=()=>{ if(!confirmLeave()) return; files=[];prepared=false;sid=null;edits.clear();dels.clear();
  $("ed").className="hide";render()};
function badge(state){
  const b=$("badge"),m={idle:"待機中",running:"実行中",success:"完了",failed:"失敗"};
  b.className="badge "+state; b.textContent=m[state]||state;
}
$("go").onclick=async()=>{
  $("go").disabled=true; $("clear").disabled=true;
  $("log").className=""; if(!prepared) $("log").textContent=""; offset=0;
  badge("running"); $("time").className="badge";
  try{
    if(prepared){
      if(edits.size+dels.size){
        if(!confirm("未保存の変更があります。保存してから取り込みますか？")) throw new Error("中止しました");
        await saveEdits();
      }
      $("ed").className="hide"; $("edit").disabled=true;
    } else await uploadAll();
    const s=await fetch(`/start?sid=${sid}`,{method:"POST"});
    if(!s.ok) throw new Error(await s.text());
    timer=setInterval(poll,1000); poll();
  }catch(e){
    $("log").textContent+="\\n❌ "+e.message+"\\n"; badge("failed");
    $("go").disabled=false; $("clear").disabled=false;
  }
};
async function poll(){
  const r=await fetch(`/log?sid=${sid}&offset=${offset}`);
  if(!r.ok) return;
  const d=await r.json();
  if(d.lines.length){
    const pre=$("log"), stick=pre.scrollTop+pre.clientHeight>=pre.scrollHeight-30;
    pre.textContent+=d.lines.join("\\n")+"\\n";
    if(stick) pre.scrollTop=pre.scrollHeight;
  }
  offset=d.offset;
  $("time").textContent=`${Math.floor(d.elapsed/60)}分${String(d.elapsed%60).padStart(2,"0")}秒`;
  if(d.state!=="running"){
    clearInterval(timer); badge(d.state); prepared=false; render();
    $("go").disabled=false; $("clear").disabled=false;
  }
}
</script></body></html>
"""


class Handler(BaseHTTPRequestHandler):
    server_version = "statscast-ingest"

    def log_message(self, fmt: str, *args) -> None:  # アクセスログは出さない
        pass

    # -- helpers ------------------------------------------------
    def _send(self, code: int, body: bytes, ctype: str) -> None:
        self.send_response(code)
        self.send_header("Content-Type", ctype)
        self.send_header("Content-Length", str(len(body)))
        self.end_headers()
        self.wfile.write(body)

    def _json(self, obj: dict, code: int = 200) -> None:
        self._send(code, json.dumps(obj).encode(), "application/json; charset=utf-8")

    def _err(self, code: int, msg: str) -> None:
        self._send(code, msg.encode(), "text/plain; charset=utf-8")

    def _query(self) -> dict[str, str]:
        q = urllib.parse.urlparse(self.path).query
        return {k: v[0] for k, v in urllib.parse.parse_qs(q).items()}

    def _job(self) -> Job | None:
        return JOBS.get(self._query().get("sid", ""))

    def _body(self) -> dict:
        length = int(self.headers.get("Content-Length", "0"))
        return json.loads(self.rfile.read(length) or b"{}") if length else {}

    def _edit_job(self) -> Job | None:
        """編集系 API 用: 準備済みかつ未開始のジョブだけ返す (それ以外はエラー送信済み)。"""
        job = self._job()
        if not job:
            self._err(404, "unknown session")
        elif not job.prepared or job.df is None:
            self._err(409, "まだ「プレビュー・編集」で準備されていません。")
        elif job.state != "idle":
            self._err(409, "取り込み開始後は編集できません。")
        else:
            return job
        return None

    # -- routes -------------------------------------------------
    def do_GET(self) -> None:
        path = urllib.parse.urlparse(self.path).path
        if path == "/":
            self._send(200, PAGE.encode(), "text/html; charset=utf-8")
        elif path == "/log":
            job = self._job()
            if not job:
                return self._err(404, "unknown session")
            try:
                offset = int(self._query().get("offset", "0"))
            except ValueError:
                offset = 0
            self._json(job.snapshot(offset))
        elif path == "/table":
            job = self._edit_job()
            if not job:
                return
            q = self._query()
            with job.edit_lock:
                idx = filtered_index(job, q)
                cols = (list(job.df.columns) if q.get("all") == "1"
                        else [c for c in KEY_COLUMNS if c in job.df.columns])
                try:
                    page, size = max(int(q.get("page", "0")), 0), min(int(q.get("size", "100")), 500)
                except ValueError:
                    page, size = 0, 100
                sel = idx[page * size:(page + 1) * size]
                sub = job.df.loc[sel, cols]
                self._json({
                    "columns": cols,
                    "allColumns": list(job.df.columns),
                    "rows": [[int(i), job.flags.get(int(i), ""), *r]
                             for i, r in zip(sub.index, sub.values.tolist())],
                    "total": len(idx),
                    "page": page, "size": size,
                    "summary": flag_summary(job),
                })
        else:
            self._err(404, "not found")

    def do_POST(self) -> None:
        global CURRENT
        path = urllib.parse.urlparse(self.path).path

        if path == "/session":
            with CURRENT_LOCK:
                if CURRENT is not None:
                    return self._err(409, "別の取り込みが実行中です。完了を待ってください。")
            base_sid = datetime.now().strftime("%Y%m%d_%H%M%S")
            sid = base_sid
            i = 2
            while sid in JOBS:  # 同一秒に 2 回押されても前のジョブを潰さない
                sid = f"{base_sid}_{i}"
                i += 1
            JOBS[sid] = Job(sid)
            return self._json({"sid": sid})

        if path == "/upload":
            job = self._job()
            if not job:
                return self._err(404, "unknown session")
            if job.state != "idle":
                return self._err(409, "この取り込みは既に開始しています。")
            name = os.path.basename(self._query().get("name", "upload.csv"))
            if not name.lower().endswith((".csv", ".xlsx", ".zip")):
                return self._err(400, f"CSV/xlsx/zip ではありません: {name}")
            length = int(self.headers.get("Content-Length", "0"))
            if length <= 0 or length > MAX_UPLOAD:
                return self._err(413, f"サイズが不正です: {length} bytes")
            dest = unique_path(job.dir, name)
            remaining = length
            with open(dest, "wb") as out:
                while remaining > 0:
                    chunk = self.rfile.read(min(1 << 20, remaining))
                    if not chunk:
                        break
                    out.write(chunk)
                    remaining -= len(chunk)
            if remaining:
                dest.unlink(missing_ok=True)
                return self._err(400, "アップロードが途中で切れました。")
            job.files.append(dest)
            return self._json({"ok": True, "saved": str(dest), "bytes": length})

        if path == "/prepare":
            job = self._job()
            if not job:
                return self._err(404, "unknown session")
            if job.state != "idle" or not job.files:
                return self._err(409, "ファイルが無いか、既に開始しています。")
            try:
                with job.edit_lock:
                    if not job.prepared:
                        job.log("━━━ 入力ファイルの準備 (編集用) ━━━")
                        csvs = collect_csvs(job)
                        rows = combine_csvs(csvs, job.dir / "combined.csv", job) if csvs else 0
                        if rows == 0:
                            return self._err(400, "投球データのファイルがありません。\n" + "\n".join(job.lines))
                        shutil.copy(job.dir / "combined.csv", job.dir / "combined_original.csv")
                        load_for_edit(job)
                        job.prepared = True
                    return self._json({"ok": True, "log": job.lines, "summary": flag_summary(job)})
            except Exception as exc:  # noqa: BLE001
                return self._err(500, f"{type(exc).__name__}: {exc}")

        if path == "/save":
            # body: {"cells": [[row, col, value], ...], "delete": [row, ...]}
            job = self._edit_job()
            if not job:
                return
            body = self._body()
            with job.edit_lock:
                df = job.df
                changed = []
                for r, c, v in body.get("cells", []):
                    if r in df.index and c in df.columns and df.at[r, c] != str(v):
                        changed.append([int(r), c, df.at[r, c], str(v)])
                        df.at[r, c] = str(v)
                dels = [int(r) for r in body.get("delete", []) if r in df.index]
                deleted = [{"row": r, "PitchNo": df.at[r, "PitchNo"] if "PitchNo" in df.columns else "",
                            "GameID": df.at[r, "GameID"] if "GameID" in df.columns else ""} for r in dels]
                job.df = df.drop(index=dels)
                if changed or dels:
                    save_edited(job, {"op": "save", "cells": changed, "deleted": deleted})
                return self._json({"ok": True, "changed": len(changed), "deleted": len(dels),
                                   "summary": flag_summary(job)})

        if path == "/replace":
            # 列の一括置換 (GameID の付け替え、日付の修正など)。現在の絞り込みに効く。
            # body: {"col", "find", "replace", "whole": bool, "fcol", "fval", "issues"}
            job = self._edit_job()
            if not job:
                return
            b = self._body()
            col = b.get("col", "")
            if col not in job.df.columns:
                return self._err(400, f"列がありません: {col}")
            find, repl = str(b.get("find", "")), str(b.get("replace", ""))
            with job.edit_lock:
                idx = filtered_index(job, {k: str(b.get(k, "")) for k in ("fcol", "fval", "issues")})
                cur = job.df.loc[idx, col]
                if b.get("whole"):
                    hit = cur == find
                    new = cur.where(~hit, repl)
                else:
                    if not find:
                        return self._err(400, "検索文字が空です (セル全体の置換は「完全一致」で)。")
                    hit = cur.str.contains(find, regex=False)
                    new = cur.str.replace(find, repl, regex=False)
                n = int(hit.sum())
                if n:
                    job.df.loc[idx, col] = new
                    save_edited(job, {"op": "replace", "col": col, "find": find,
                                      "replace": repl, "whole": bool(b.get("whole")), "rows": n})
                return self._json({"ok": True, "replaced": n, "summary": flag_summary(job)})

        if path == "/revert":
            job = self._edit_job()
            if not job:
                return
            with job.edit_lock:
                shutil.copy(job.dir / "combined_original.csv", job.dir / "combined.csv")
                load_for_edit(job)
                with open(job.dir / "edits.jsonl", "a", encoding="utf-8") as fh:
                    fh.write(json.dumps({"at": datetime.now().isoformat(timespec="seconds"),
                                         "op": "revert"}) + "\n")
            return self._json({"ok": True, "summary": flag_summary(job)})

        if path == "/start":
            job = self._job()
            if not job:
                return self._err(404, "unknown session")
            if not job.files:
                return self._err(400, "ファイルがありません。")
            with CURRENT_LOCK:
                if CURRENT is not None:
                    return self._err(409, "別の取り込みが実行中です。")
                CURRENT = job
            threading.Thread(target=run_job, args=(job,), daemon=True).start()
            return self._json({"ok": True})

        self._err(404, "not found")


def main() -> int:
    if not DEPLOY_SH.exists():
        print(f"エラー: {DEPLOY_SH} が見つかりません。", file=sys.stderr)
        return 1
    INBOX.mkdir(parents=True, exist_ok=True)
    srv = ThreadingHTTPServer((HOST, PORT), Handler)
    url = f"http://{HOST}:{PORT}"
    print("=" * 52)
    print(" statscast 取込アプリ")
    print(f"   {url}  をブラウザで開いてください")
    print(f"   受信ファイル置き場: {INBOX}")
    print("   終了: Ctrl-C")
    print("=" * 52)
    try:
        srv.serve_forever()
    except KeyboardInterrupt:
        print("\n終了しました。")
    return 0


if __name__ == "__main__":
    sys.exit(main())
